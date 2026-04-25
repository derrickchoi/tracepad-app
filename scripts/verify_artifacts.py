from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

required = [
    ROOT / "apps/web/dist/index.html",
    ROOT / "financials/tracepad-five-year-model.xlsx",
    ROOT / "deck/tracepad-seed-deck.pptx",
    ROOT / "docs/product/tracepad-v1-prd.docx",
    ROOT / "brand/tracepad-logo-lockup.png",
]

missing = [str(p) for p in required if not p.exists()]
if missing:
    raise SystemExit("Missing required artifacts:\n" + "\n".join(missing))

wb = load_workbook(ROOT / "financials/tracepad-five-year-model.xlsx", data_only=False)
required_sheets = {"Executive Summary", "Assumptions", "Customer Growth", "Revenue Model", "Pricing & Packaging", "COGS Service Costs", "Headcount Plan", "OpEx", "P&L", "Cash Runway", "Scenarios", "Charts"}
if not required_sheets.issubset(set(wb.sheetnames)):
    raise SystemExit("Workbook missing required sheets")
formula_errors = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(err in cell.value for err in ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"]):
                formula_errors.append(f"{ws.title}!{cell.coordinate}")
if formula_errors:
    raise SystemExit("Formula error tokens found: " + ", ".join(formula_errors))

with ZipFile(ROOT / "deck/tracepad-seed-deck.pptx") as zf:
    slides = [name for name in zf.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
    if len(slides) < 10:
        raise SystemExit("Pitch deck has fewer than 10 slides")

print("Artifact verification passed")
